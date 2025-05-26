
from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from PyPDF2 import PdfReader
import docx
import openpyxl

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.post("/summarize/")
async def summarize(file: UploadFile = File(...)):
    content = await file.read()
    summary = "Unsupported file type"
    if file.filename.endswith(".txt"):
        summary = content.decode("utf-8")[:100]
    elif file.filename.endswith(".pdf"):
        reader = PdfReader(file.file)
        text = "".join([page.extract_text() for page in reader.pages])
        summary = text[:500]
    elif file.filename.endswith(".docx"):
        doc = docx.Document(file.file)
        text = "
".join([para.text for para in doc.paragraphs])
        summary = text[:500]
    elif file.filename.endswith(".xlsx"):
        wb = openpyxl.load_workbook(file.file)
        text = ""
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                text += " ".join([str(cell) for cell in row if cell]) + "\n"
        summary = text[:500]
    return {"summary": summary}
